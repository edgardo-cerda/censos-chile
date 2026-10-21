"""Planilla con las divisiones territoriales usadas en cada censo (1813-1982) y su comparabilidad.

Lee las unidades territoriales de los cuadros ya transcritos (productos/tabulados/
tabulados_publicaciones.csv y, para 1843-1865, los CSV de digitalizacion/1865) y escribe
productos/territorio/divisiones_territoriales_censos.xlsx con las hojas:
  - Léeme: qué hay en cada hoja, fuentes y advertencias.
  - Resumen: una fila por censo; los conteos son fórmulas sobre la hoja Unidades.
  - Provincias: matriz provincia × censo (1835-1970) con el estado de cada provincia frente al censo
    anterior según sus departamentos (igual, cambia, nueva); al lado, la composición.
  - Departamentos: matriz departamento × censo (1835-1952) con la provincia de cada censo.
  - Comunas: matriz comuna × censo (1930, 1940, 1952 y 1982) por nombre.
  - Cambios: cada creación, supresión o traspaso entre censos consecutivos.
  - Comprobaciones: cifras de los cuadros comparativos (que presentan censos anteriores) frente a
    la cifra original de cada censo; si son iguales, la unidad no cambió entre ambos censos.
  - Series comparables: resumen de esos cuadros comparativos.
  - Unidades: lista larga de todas las unidades de cada censo, con su unidad superior y población.

Las unidades se igualan entre censos por su nombre normalizado (ALIAS); un mismo nombre no
asegura los mismos límites: eso solo lo prueban las comprobaciones.

Uso:
  python herramientas/divisiones_territoriales.py
"""

from __future__ import annotations

import unicodedata
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
TABULADOS = RAIZ / "productos" / "tabulados" / "tabulados_publicaciones.csv"
SALIDA = RAIZ / "productos" / "territorio" / "divisiones_territoriales_censos.xlsx"

# Cuadro de la propia publicación del que sale la división de cada censo (fuente_id, cuadro).
PROPIO = {
    1813: ("pub_1813", "censo_provincia_distritos"), 1835: ("pub_1835", "art34_42"),
    1875: ("pub_1875", "602"), 1885: ("pub_1885", "839"), 1895: ("pub_1895", "IV.467"),
    1907: ("pub_1907", "1"), 1920: ("pub_1920", "1"), 1930: ("pub_1930", "9"), 1940: ("pub_1940", "p76"),
    1952: ("pub_1952", "N4"), 1960: ("pub_1960", "1.1.1"), 1970: ("pub_1970", "5"), 1982: ("pub_1982_v1", "1.01"),
}
CUADRO_TEXTO = {
    1813: "Censo de 1813, partidos y distritos (PDF 62 y sigs.)",
    1835: "Censo de 1835, arts. 34-42 (departamentos por provincia)",
    1843: "Sinóptico del censo de 1865 (pág. 392): departamentos y provincias de 1843",
    1854: "Sinóptico del censo de 1865 (pág. 392): departamentos y provincias de 1854",
    1865: "Censo de 1865, págs. 352 (departamentos) y 360 (provincias)",
    1875: "Censo de 1875, cuadro general (págs. 602-605)",
    1885: "Censo de 1885, cuadro general (tomo I, pág. 839)",
    1895: "Censo de 1895, cuadro general (tomo IV, pág. 467)",
    1907: "Censo de 1907, cuadro 1",
    1920: "Censo de 1920, cuadro 1",
    1930: "Censo de 1930, cuadro 9 (provincias, departamentos y comunas)",
    1940: "Censo de 1940, pág. 76 (provincias, departamentos y comunas)",
    1952: "Censo de 1952, cuadro nacional 4 (provincias, departamentos y comunas)",
    1960: "Censo de 1960, cuadro 1.1.1 (provincias)",
    1970: "Censo de 1970, cuadro 5 (provincias)",
    1982: "Censo de 1982, cuadro 1.01 (regiones, provincias y comunas; DPA de 1974)",
}
CENSOS = [1813, 1835, 1843, 1854, 1865, 1875, 1885, 1895, 1907, 1920, 1930, 1940, 1952, 1960, 1970, 1982]
C_PROV = [1835, 1843, 1854, 1865, 1875, 1885, 1895, 1907, 1920, 1930, 1940, 1952, 1960, 1970]
C_DEPTO = [1835, 1843, 1854, 1865, 1875, 1885, 1895, 1907, 1920, 1930, 1940, 1952]
C_COMUNA = [1930, 1940, 1952, 1982]

# Variantes de nombre de una misma unidad (clave normalizada -> clave canónica).
ALIAS = {
    "provincia": {"aysen": "aisen", "antartida": "antartica", "antartico": "antartica",
                  "territorio antartico": "antartica"},
    "departamento": {"laja": "la laja", "union": "la union", "serena": "la serena", "ligua": "la ligua",
                     "andes": "los andes", "puchacai": "puchacay", "loa": "el loa",
                     "nueva imperial": "imperial", "aysen": "aisen", "quirihue": "itata"},
    "comuna": {"llayllay": "llaillay", "puerto natales": "natales", "ligua": "la ligua", "puerto saavedra": "saavedra",
               "san jose de la mariquina": "mariquina", "aysen": "aisen", "cisterna": "la cisterna",
               "gral. lagos": "general lagos", "tilcoco": "quinta de tilcoco", "coihayque": "coihaique",
               "paihuano": "paiguano", "marchigue": "marchihue", "rio baker": "baker"},
}
# Alias que dependen de la provincia: Victoria de Santiago (después La Victoria) no es la de
# Malleco; San Carlos de Chiloé (1835) es San Carlos de Ancud, no el de Ñuble.
ALIAS_CONTEXTO = {("santiago", "victoria"): "la victoria", ("chiloe", "san carlos"): "san carlos de ancud"}
NOMBRE = {"san carlos de ancud": "San Carlos (de Ancud)", "antartica": "Antártica", "aisen": "Aisén"}
PREFIJOS = ("territorio de colonizacion de ", "territorio de ", "colonia de ", "t. de ", "t. ")

ARIAL = "Arial"
F_NORMAL = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITULO = Font(name=ARIAL, size=14, bold=True)
F_NOTA = Font(name=ARIAL, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
COLOR = {"igual": "C6EFCE", "cambia": "FFEB9C", "nueva": "BDD7EE", "traslado": "FFEB9C",
         "sin detalle": "EDEDED", "distinta": "FFC7CE", "existe": "E2EFDA", "vuelve": "DDEBF7", "errata": "F8CBAD"}


def sin_tildes(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def clave(nombre: str, nivel: str, provincia: str = "") -> str:
    k = " ".join(sin_tildes(nombre).lower().replace("’", "'").split())
    if nivel == "comuna":
        k = k.replace("-", "")
    for p in PREFIJOS:
        if k.startswith(p):
            k = k[len(p):]
    k = ALIAS_CONTEXTO.get((provincia, k), k) if nivel == "departamento" else k
    return ALIAS.get(nivel, {}).get(k, k)


def unidades() -> pd.DataFrame:
    """Una fila por unidad y censo: censo, nivel, nombre, clave, provincia, departamento, región, población."""
    d = pd.read_csv(TABULADOS, dtype=str, keep_default_na=False)
    filas = []
    for censo, (fuente, cuadro) in PROPIO.items():
        x = d[(d.censo == str(censo)) & (d.fuente_id == fuente) & (d.cuadro == cuadro) & (d.area == "_T")
              & (d.edad_min == "") & (d.territorio_nivel != "pais")]
        if censo == 1813:  # solo por sexo: se suman hombres y mujeres
            x = x[x.sexo.isin(["hombre", "mujer"])].copy()
            x["valor"] = pd.to_numeric(x.valor, errors="coerce")
            x = x.groupby(["territorio_cod_orig", "territorio_nivel", "territorio_nombre_orig"], sort=False,
                          as_index=False).valor.sum()
        else:
            x = x[x.sexo == "_T"].drop_duplicates("territorio_cod_orig")
        nombres = dict(zip(x.territorio_cod_orig, x.territorio_nombre_orig))
        niveles = dict(zip(x.territorio_cod_orig, x.territorio_nivel))
        for _, f in x.iterrows():
            partes = f.territorio_cod_orig.split("/")
            sup = ["/".join(partes[:i]) for i in range(1, len(partes))]
            arriba = {niveles[s]: nombres[s] for s in sup if s in niveles}
            filas.append({"censo": censo, "nivel": f.territorio_nivel, "nombre": f.territorio_nombre_orig,
                          "region": arriba.get("region", ""), "provincia": arriba.get("provincia", ""),
                          "departamento": arriba.get("departamento", ""), "partido": arriba.get("partido", ""),
                          "poblacion": pd.to_numeric(f.valor, errors="coerce"), "fuente": f"{fuente} c{cuadro}",
                          "estado": f.get("estado_valor", "")})
    # 1865: departamentos con su provincia (pág. 352) y provincias (pág. 360)
    dig = RAIZ / "digitalizacion" / "1865"
    p352 = pd.read_csv(dig / "pub_1865_p352.csv", dtype=str)
    for _, f in p352[p352.nivel == "departamento"].iterrows():
        filas.append({"censo": 1865, "nivel": "departamento", "nombre": f.territorio, "provincia": f.provincia,
                      "poblacion": int(f.total_general), "fuente": "pub_1865 cp352"})
    p360 = pd.read_csv(dig / "pub_1865_p360.csv", dtype=str)
    for _, f in p360[p360.nivel == "provincia"].iterrows():
        filas.append({"censo": 1865, "nivel": "provincia", "nombre": f.territorio,
                      "poblacion": int(f.total_general), "fuente": "pub_1865 cp360"})
    # 1843 y 1854: sinóptico de 1865. La provincia de cada departamento es la de 1865, salvo las
    # que en 1843 aún no existían (composición según las llaves del original, ver pub_1865_p392_prov).
    p392 = pd.read_csv(dig / "pub_1865_p392.csv", dtype=str)
    prov = pd.read_csv(dig / "pub_1865_p392_prov.csv", dtype=str)
    en_1843 = {"Carelmapu": "Chiloé", "Osorno": "Valdivia", "Laja": "Concepción", "Chillán": "Concepción",
               "San Cárlos": "Maule", "Copiapó": "Coquimbo", "Vallenar": "Coquimbo", "Freirina": "Coquimbo"}
    for censo in (1843, 1854):
        for _, f in p392[(p392.nivel == "departamento") & (p392[str(censo)] != "..")].iterrows():
            pr = en_1843.get(f.territorio, f.provincia_1865) if censo == 1843 else f.provincia_1865
            filas.append({"censo": censo, "nivel": "departamento", "nombre": f.territorio, "provincia": pr,
                          "poblacion": int(f[str(censo)]), "fuente": "pub_1865 cp392"})
        for _, f in prov[(prov.censo == str(censo)) & (prov.territorio != "Total")].iterrows():
            filas.append({"censo": censo, "nivel": "provincia", "nombre": f.territorio,
                          "poblacion": int(f.poblacion), "fuente": "pub_1865 cp392"})
    u = pd.DataFrame(filas).fillna("")
    u["clave_prov"] = [clave(p, "provincia") if p else "" for p in u.provincia]
    u["clave"] = [clave(n, lv, cp) for n, lv, cp in zip(u.nombre, u.nivel, u.clave_prov)]
    u.loc[u.nivel == "provincia", "clave_prov"] = u.clave
    u["clave_depto"] = [clave(dp, "departamento", cp) if dp else "" for dp, cp in zip(u.departamento, u.clave_prov)]
    u.loc[u.nivel == "departamento", "clave_depto"] = u.clave
    # 1843 y 1854: la suma de los departamentos debe dar el total de cada provincia del sinóptico
    for censo in (1843, 1854):
        x = u[u.censo == censo]
        suma = x[x.nivel == "departamento"].groupby("clave_prov").poblacion.sum()
        tot = x[x.nivel == "provincia"].set_index("clave").poblacion
        assert (suma.reindex(tot.index).fillna(0) == tot).all(), (censo, suma, tot)
    return u


def nombre_visible(u: pd.DataFrame, nivel: str, k: str) -> str:
    """Nombre de la unidad tal como se imprimió en el censo más reciente (salvo 1982, en mayúsculas)."""
    if k in NOMBRE:
        return NOMBRE[k]
    x = u[(u.nivel == nivel) & (u.clave == k)].sort_values("censo")
    no82 = x[x.censo != 1982]
    n = (no82 if len(no82) else x).nombre.iloc[-1]
    return n.title() if n.isupper() else n


def cabecera(ws, fila, titulos, anchos):
    for j, (t, w) in enumerate(zip(titulos, anchos), 1):
        c = ws.cell(fila, j, t)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if w:
            ws.column_dimensions[get_column_letter(j)].width = w


def celda(ws, r, c, v, estado=None, font=F_NORMAL, wrap=False, centro=False):
    x = ws.cell(r, c, v)
    x.font = font
    if estado:
        x.fill = PatternFill("solid", fgColor=COLOR[estado])
    if wrap or centro:
        x.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal="center" if centro else None)
    return x


def lista(nombres) -> str:
    return ", ".join(sorted(nombres))


def main() -> None:
    u = unidades()
    wb = Workbook()
    cambios = []  # (censo anterior, censo, nivel, unidad, tipo, detalle)

    # ---- Provincias ----
    prov = u[(u.nivel == "provincia") & u.censo.isin(C_PROV)]
    deps = u[u.nivel == "departamento"]
    comp = {(c, k): set(deps[(deps.censo == c) & (deps.clave_prov == k)].clave) for c, k in zip(prov.censo, prov.clave)}
    orden_p = list(dict.fromkeys(prov.sort_values("censo", ascending=False).clave))
    # orden: el de la última aparición, de norte a sur, según la secuencia impresa en cada censo
    orden_p = sorted(orden_p, key=lambda k: (-prov[prov.clave == k].censo.max(),
                                             list(prov[prov.censo == prov[prov.clave == k].censo.max()].clave).index(k)))
    nom_d = {k: nombre_visible(u, "departamento", k) for k in deps.clave.unique()}
    wp = wb.active
    wp.title = "Provincias"
    titulos = ["Provincia"] + [str(c) for c in C_PROV] + ["Censos", "Tramos sin cambio de departamentos"] + \
              [f"Departamentos {c}" for c in C_PROV if c <= 1952]
    cabecera(wp, 1, titulos, [22] + [9] * len(C_PROV) + [8, 34] + [34] * len([c for c in C_PROV if c <= 1952]))
    for i, k in enumerate(orden_p, 2):
        celda(wp, i, 1, nombre_visible(u, "provincia", k), font=F_BOLD)
        tramos, actual, previo = [], None, None
        for j, c in enumerate(C_PROV):
            col = j + 2
            if (c, k) not in comp:
                if previo is not None and c != 1982:
                    cambios.append((C_PROV[j - 1], c, "provincia", nombre_visible(u, "provincia", k),
                                    "deja de aparecer", lista(nom_d[x] for x in comp.get((C_PROV[j - 1], k), set()))))
                if actual:
                    tramos.append(actual)
                actual, previo = None, None
                continue
            ds = comp[(c, k)]
            ant = C_PROV[j - 1] if j else None
            if previo is None:
                antes = any((x, k) in comp for x in C_PROV[:j])
                est = ("vuelve" if antes else "nueva") if j else "existe"
                if j:
                    cambios.append((ant, c, "provincia", nombre_visible(u, "provincia", k),
                                    "vuelve a aparecer" if antes else "nueva",
                                    lista(nom_d[x] for x in ds) if ds else ""))
            elif ds == previo and c <= 1952:
                est = "igual"
            elif not ds or not previo:
                est = "sin detalle"
            else:
                est = "cambia"
                det = []
                if ds - previo:
                    det.append("entra " + lista(nom_d[x] for x in ds - previo))
                if previo - ds:
                    det.append("sale " + lista(nom_d[x] for x in previo - ds))
                cambios.append((ant, c, "provincia", nombre_visible(u, "provincia", k), "cambia departamentos",
                                "; ".join(det)))
            texto = {"igual": "=", "cambia": "Δ", "nueva": "nueva", "vuelve": "vuelve", "existe": "✓",
                     "sin detalle": "✓"}[est]
            celda(wp, i, col, texto, est, centro=True)
            if est in ("igual", "sin detalle") and actual:
                actual[1] = c
            else:
                if actual:
                    tramos.append(actual)
                actual = [c, c]
            previo = ds
        if actual:
            tramos.append(actual)
        n = len(C_PROV)
        celda(wp, i, n + 2, f'=COUNTIF(B{i}:{get_column_letter(n + 1)}{i},"?*")', centro=True)
        celda(wp, i, n + 3, ", ".join(f"{a}–{b}" if a != b else str(a) for a, b in tramos), wrap=True)
        for j, c in enumerate([c for c in C_PROV if c <= 1952]):
            if (c, k) in comp:
                ds = comp[(c, k)]
                celda(wp, i, n + 4 + j, lista(nom_d[x] for x in ds) if ds else "(sin departamentos)", wrap=True)
    wp.freeze_panes = "B2"
    wp.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{len(orden_p) + 1}"
    nota_p = len(orden_p) + 3
    for t in ["Signos: «=» mismos departamentos que en el censo anterior · Δ cambian los departamentos (ver Cambios) · "
              "nueva: no estaba en el censo anterior · ✓ existe, sin detalle de departamentos (1960 y 1970) o primer censo.",
              "Tramos sin cambio: censos seguidos en que la provincia tuvo los mismos departamentos (por nombre). En 1960 y "
              "1970 solo se sabe que la provincia existe; el tramo se extiende por nombre.",
              "1843 y 1854 vienen del sinóptico de 1865, que agrupa los departamentos como en 1865 salvo las provincias que "
              "aún no existían. 1982 usa otra división (regiones y provincias de 1974): ver la hoja Comunas y Unidades."]:
        celda(wp, nota_p, 1, t, font=F_NOTA)
        nota_p += 1

    # ---- Departamentos ----
    wd = wb.create_sheet("Departamentos")
    dd = deps[deps.censo.isin(C_DEPTO)]
    prov_de = {(c, k): p for c, k, p in zip(dd.censo, dd.clave, dd.clave_prov)}
    nom_p = {k: nombre_visible(u, "provincia", k) for k in u[u.nivel == "provincia"].clave.unique()}
    nom_p.update({k: k.title() for k in dd.clave_prov.unique() if k not in nom_p})
    orden_d = sorted(dd.clave.unique(), key=lambda k: (
        orden_p.index(prov_de[(dd[dd.clave == k].censo.max(), k)]) if prov_de[(dd[dd.clave == k].censo.max(), k)] in orden_p else 99,
        -dd[dd.clave == k].censo.max(), nom_d[k]))
    titulos = ["Departamento", "Provincia (último censo)"] + [str(c) for c in C_DEPTO] + ["Censos", "Cambios de provincia"]
    cabecera(wd, 1, titulos, [24, 20] + [13] * len(C_DEPTO) + [8, 40])
    for i, k in enumerate(orden_d, 2):
        ult = dd[dd.clave == k].censo.max()
        celda(wd, i, 1, nom_d[k], font=F_BOLD)
        celda(wd, i, 2, nom_p[prov_de[(ult, k)]])
        previo, traslados = None, []
        for j, c in enumerate(C_DEPTO):
            p = prov_de.get((c, k))
            if p is None:
                if previo is not None:
                    cambios.append((C_DEPTO[j - 1], c, "departamento", nom_d[k], "deja de aparecer",
                                    f"estaba en {nom_p[previo]}"))
                previo = None
                continue
            if previo is None:
                antes = any((x, k) in prov_de for x in C_DEPTO[:j])
                est = ("vuelve" if antes else "nueva") if j else "existe"
                if j:
                    cambios.append((C_DEPTO[j - 1], c, "departamento", nom_d[k],
                                    "vuelve a aparecer" if antes else "nuevo", f"en {nom_p[p]}"))
            elif p == previo:
                est = "igual"
            else:
                est = "traslado"
                traslados.append(f"{c}: {nom_p[previo]} → {nom_p[p]}")
                cambios.append((C_DEPTO[j - 1], c, "departamento", nom_d[k], "cambia de provincia",
                                f"{nom_p[previo]} → {nom_p[p]}"))
            celda(wd, i, j + 3, nom_p[p], est)
            previo = p
        n = len(C_DEPTO)
        celda(wd, i, n + 3, f'=COUNTIF(C{i}:{get_column_letter(n + 2)}{i},"?*")', centro=True)
        celda(wd, i, n + 4, "; ".join(traslados), wrap=True)
    wd.freeze_panes = "C2"
    wd.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{len(orden_d) + 1}"
    r = len(orden_d) + 3
    for t in ["Cada celda indica la provincia del departamento en ese censo. Verde: misma provincia que en el censo "
              "anterior · amarillo: cambió de provincia · azul: no estaba en el censo anterior.",
              "Mismo nombre no implica mismos límites (p. ej. Itata o Laja en 1930). Para saber si un departamento no "
              "cambió, ver la hoja Comprobaciones (cuadro 1 de 1920 para 1865-1907).",
              "1960 y 1970 no tienen departamentos transcritos; en 1982 ya no existen (DPA de 1974)."]:
        celda(wd, r, 1, t, font=F_NOTA)
        r += 1

    # ---- Comunas ----
    wc = wb.create_sheet("Comunas")
    cm = u[(u.nivel == "comuna") & u.censo.isin(C_COMUNA)]
    sup = {(c, k): (dep if c != 1982 else p) for c, k, dep, p in zip(cm.censo, cm.clave, cm.departamento, cm.provincia)}
    nom_c = {k: nombre_visible(u, "comuna", k) for k in cm.clave.unique()}
    reg82 = dict(zip(cm[cm.censo == 1982].clave, cm[cm.censo == 1982].region))
    orden_c = sorted(cm.clave.unique(), key=lambda k: (reg82.get(k) is None, nom_c[k]))
    orden_c = list(cm[cm.censo == 1982].clave.drop_duplicates()) + sorted(
        set(cm.clave) - set(cm[cm.censo == 1982].clave), key=lambda k: nom_c[k])
    titulos = ["Comuna", "Región 1982"] + [f"{c} ({'provincia' if c == 1982 else 'departamento'})" for c in C_COMUNA] + \
              ["Censos", "Con el mismo nombre"]
    cabecera(wc, 1, titulos, [26, 22] + [20] * len(C_COMUNA) + [8, 22])
    for i, k in enumerate(orden_c, 2):
        celda(wc, i, 1, nom_c[k], font=F_BOLD)
        r82 = reg82.get(k, "")
        celda(wc, i, 2, r82.title() if r82.isupper() else r82)
        presentes = []
        previo = False
        for j, c in enumerate(C_COMUNA):
            s = sup.get((c, k))
            if s is None:
                if previo:
                    cambios.append((C_COMUNA[j - 1], c, "comuna", nom_c[k], "deja de aparecer", ""))
                previo = False
                continue
            presentes.append(c)
            est = "existe" if previo or j == 0 else "nueva"
            if j and not previo:
                cambios.append((C_COMUNA[j - 1], c, "comuna", nom_c[k], "nueva", f"en {s.title() if s.isupper() else s}"))
            celda(wc, i, j + 3, s.title() if s.isupper() else s, est)
            previo = True
        n = len(C_COMUNA)
        celda(wc, i, n + 3, f'=COUNTIF(C{i}:{get_column_letter(n + 2)}{i},"?*")', centro=True)
        celda(wc, i, n + 4, f"{presentes[0]}–{presentes[-1]}" if len(presentes) > 1 and presentes == C_COMUNA[
            C_COMUNA.index(presentes[0]):C_COMUNA.index(presentes[-1]) + 1] else ", ".join(map(str, presentes)))
    wc.freeze_panes = "C2"
    wc.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{len(orden_c) + 1}"
    r = len(orden_c) + 3
    for t in ["Comunas igualadas solo por nombre (sin tildes ni mayúsculas, y con variantes de escritura como Til-Til y "
              "Tiltil, Coihayque y Coihaique, Paihuano y Paiguano, Puerto Natales y Natales, Ligua y La Ligua). Entre 1930 y 1952 la celda muestra el "
              "departamento; en 1982, la provincia de la DPA de 1974. Verde: también estaba en el censo anterior · "
              "azul: nombre nuevo.",
              "Una comuna con el mismo nombre puede haber cambiado de límites (divisiones y fusiones de 1970-1981); la "
              "comparación exacta 1970-1982 está en el volumen de cifras comparativas (copia de Memoria Chilena, "
              "sin transcribir)."]:
        celda(wc, r, 1, t, font=F_NOTA)
        r += 1

    # ---- Comprobaciones con los cuadros comparativos ----
    d = pd.read_csv(TABULADOS, dtype=str, keep_default_na=False)
    orig = {(int(c), lv, k): v for c, lv, k, v in zip(u.censo, u.nivel, u.clave, u.poblacion)}
    comparativos = [
        ("pub_1865", "p392", "Sinóptico de 1865 (pág. 392)", [1835], "la de cada censo"),
        ("pub_1875", "p645", "Sinóptico de 1875 (pág. 645)", [1835, 1865], "la de cada censo"),
        ("pub_1920", "1", "Cuadro 1 de 1920", [1865, 1875, 1885, 1895, 1907], "la de 1920"),
        ("pub_1930", "10", "Cuadro 10 de 1930", [1895, 1907, 1920], "la de 1930"),
        ("pub_1952", "9", "Cuadro 9 de 1952", [1865, 1875, 1885, 1895, 1907, 1920, 1930, 1940], "la de cada censo"),
        ("pub_1960", "1.1.1", "Cuadro 1.1.1 de 1960", [1940, 1952], "la de 1960"),
    ]
    SINOPTICO = {f: pd.read_csv(RAIZ / ruta, dtype=str).set_index("territorio")[col].to_dict() for f, ruta, col in (
        ("pub_1865", "digitalizacion/1865/pub_1865_p392.csv", "provincia_1865"),
        ("pub_1875", "digitalizacion/1875/pub_1875_p645.csv", "provincia_1875"))}
    est_orig = {(int(c), lv, k): e for c, lv, k, e in zip(u.censo, u.nivel, u.clave, u.estado)}
    comp_filas, resumen_comp = [], []
    for fuente, cuadro, titulo, censos, division in comparativos:
        x = d[(d.fuente_id == fuente) & (d.cuadro == cuadro) & (d.sexo == "_T") & (d.area == "_T")
              & d.territorio_nivel.isin(["provincia", "departamento"]) & d.censo.isin([str(c) for c in censos])]
        x = x.drop_duplicates(["censo", "territorio_nivel", "territorio_cod_orig"])
        for nivel in ("provincia", "departamento"):
            y = x[x.territorio_nivel == nivel]
            if not len(y):
                continue
            iguales = distintas = sin = erratas = 0
            for _, f in y.iterrows():
                v = pd.to_numeric(f.valor, errors="coerce")
                if pd.isna(v):
                    continue
                padre = f.territorio_cod_orig.split("/")[0] if "/" in f.territorio_cod_orig else ""
                cp = clave(x[x.territorio_cod_orig == padre].territorio_nombre_orig.iloc[0], "provincia") \
                    if padre and (x.territorio_cod_orig == padre).any() else ""
                if fuente in SINOPTICO and nivel == "departamento":
                    cp = clave(SINOPTICO[fuente].get(f.territorio_nombre_orig, ""), "provincia")
                k = clave(f.territorio_nombre_orig, nivel, cp)
                o = orig.get((int(f.censo), nivel, k))
                if o is None or pd.isna(o):
                    est, sin = "sin unidad del mismo nombre", sin + 1
                elif int(o) == int(v):
                    est, iguales = "igual", iguales + 1
                elif "inconsistente" in (f.estado_valor, est_orig.get((int(f.censo), nivel, k))):
                    est, erratas = "errata anotada", erratas + 1
                else:
                    est, distintas = "distinta", distintas + 1
                comp_filas.append((titulo, division, int(f.censo), nivel, f.territorio_nombre_orig, int(v),
                                   "" if o is None or pd.isna(o) else int(o),
                                   "" if o is None or pd.isna(o) else int(v) - int(o), est))
            resumen_comp.append((titulo, nivel, ", ".join(map(str, censos)), division, iguales, distintas, erratas, sin))
    wq = wb.create_sheet("Comprobaciones")
    titulos = ["Cuadro comparativo", "División que usa", "Censo", "Nivel", "Unidad (como en el comparativo)",
               "Cifra del comparativo", "Cifra original del censo", "Diferencia", "Resultado"]
    cabecera(wq, 1, titulos, [24, 16, 8, 13, 30, 14, 14, 12, 24])
    for i, f in enumerate(comp_filas, 2):
        for j, v in enumerate(f, 1):
            c = celda(wq, i, j, v)
            if j in (6, 7, 8):
                c.number_format = "#,##0;-#,##0;0"
        estado = {"igual": "igual", "distinta": "distinta", "errata anotada": "errata"}.get(f[-1])
        if estado:
            wq.cell(i, 9).fill = PatternFill("solid", fgColor=COLOR[estado])
    wq.freeze_panes = "A2"
    wq.auto_filter.ref = f"A1:I{len(comp_filas) + 1}"

    # ---- Series comparables ----
    ws = wb.create_sheet("Series comparables")
    celda(ws, 1, 1, "Cuadros que presentan cifras de censos anteriores", font=F_TITULO)
    titulos = ["Cuadro comparativo", "Nivel", "Censos anteriores", "División que usa", "Cifras iguales al original",
               "Cifras distintas", "Distintas por errata anotada", "Sin unidad del mismo nombre", "Lectura"]
    cabecera(ws, 3, titulos, [24, 13, 26, 16, 12, 10, 11, 12, 90])
    lectura = {
        ("Cuadro 1 de 1920", "departamento"): "Reconstruye 1865-1907 en los departamentos de 1920. Una cifra igual a la "
        "original indica que el departamento tenía los mismos límites en ese censo y en 1920; una distinta, que cambió "
        "(p. ej. Talca 1865 = Talca + Curepto de 1920).",
        ("Cuadro 1 de 1920", "provincia"): "Provincias de 1920 reconstruidas hacia atrás: serie comparable 1865-1920.",
        ("Cuadro 10 de 1930", "provincia"): "Reconstruye 1895-1920 en las 17 provincias de 1930 (reforma de 1928): serie "
        "comparable 1895-1930. Las cifras distintas son provincias que cambiaron con la reforma.",
        ("Cuadro 9 de 1952", "provincia"): "No reconstruye: pone cada censo con sus propias provincias, en la fila de la "
        "provincia de 1952 de igual nombre (Angol 1875-1885 en Malleco; Aconcagua 1930 incluye Valparaíso). Sirve para la "
        "serie de cada nombre, no para límites fijos.",
        ("Cuadro 1.1.1 de 1960", "provincia"): "1940 y 1952 en las provincias de 1960. Todas iguales (salvo 1940, que usa "
        "las cifras revisadas): las 25 provincias no cambiaron entre 1940 y 1960.",
        ("Sinóptico de 1865 (pág. 392)", "departamento"): "Departamentos de 1835, 1843 y 1854. Para 1835 agrupa los 50 "
        "departamentos del censo en las unidades del sinóptico; solo coinciden los que no cambiaron.",
        ("Sinóptico de 1865 (pág. 392)", "provincia"): "Provincias de 1835 con las llaves del original.",
        ("Sinóptico de 1875 (pág. 645)", "departamento"): "Mismo contenido que el de 1865 para 1835-1865 (igual cifra a "
        "cifra salvo erratas).",
        ("Sinóptico de 1875 (pág. 645)", "provincia"): "Provincias de 1835 y 1865.",
    }
    for i, f in enumerate(resumen_comp, 4):
        for j, v in enumerate(f, 1):
            celda(ws, i, j, v, wrap=j == 3)
        celda(ws, i, 9, lectura.get((f[0], f[1]), ""), wrap=True)
    r = len(resumen_comp) + 6
    celda(ws, r, 1, "Series comparables que se pueden armar", font=F_BOLD)
    series = [
        ("1835-1865", "Departamentos de cada censo (sinópticos de 1865 y 1875); límites del censo respectivo."),
        ("1865-1920", "Departamentos y provincias en la división de 1920 (cuadro 1 de 1920), con Tacna y Tarapacá desde 1885."),
        ("1895-1930", "Provincias en la división de 1930 (cuadro 10 de 1930), por área urbana y rural."),
        ("1865-1952", "Provincias por nombre (cuadros 5, 6 y 9 de 1952): cada censo con sus propias provincias; los "
         "límites cambian (sobre todo en 1928 y entre 1930 y 1940)."),
        ("1940-1960", "Las mismas 25 provincias (cuadro 1.1.1 de 1960); 1970 tiene las mismas 25 por nombre."),
        ("1970-1982", "Comunas en la DPA de 1974 (volumen de cifras comparativas 1970-1982, copia de Memoria Chilena, "
         "por transcribir)."),
        ("1982 en adelante", "Regiones, provincias y comunas de la DPA de 1974; los microdatos 1992-2024 se llevan a la "
         "DPA de 2018 (ver metadatos/variables/region.yaml)."),
    ]
    cabecera(ws, r + 1, ["Período", "Base"], [None, None])
    for i, (a, b) in enumerate(series, r + 2):
        celda(ws, i, 1, a, font=F_BOLD)
        celda(ws, i, 2, b)
    ws.freeze_panes = "A4"

    # ---- Cambios ----
    wk = wb.create_sheet("Cambios")
    titulos = ["Censo anterior", "Censo", "Nivel", "Unidad", "Cambio", "Detalle"]
    cabecera(wk, 1, titulos, [10, 8, 14, 26, 22, 80])
    orden_niv = {"provincia": 0, "departamento": 1, "comuna": 2}
    cambios.append((1835, 1843, "departamento", "Itata", "cambia de nombre",
                    "Quirihue en 1835 (Maule); el sinóptico de 1865 le da a Itata la misma cifra de 1835 (18.837)."))
    cambios.sort(key=lambda f: (f[1], orden_niv[f[2]], f[4], f[3]))
    color_c = {"nueva": "nueva", "nuevo": "nueva", "cambia departamentos": "cambia", "cambia de provincia": "traslado",
               "deja de aparecer": "sin detalle", "vuelve a aparecer": "vuelve", "cambia de nombre": "cambia"}
    for i, f in enumerate(cambios, 2):
        for j, v in enumerate(f, 1):
            celda(wk, i, j, v, wrap=j == 6)
        wk.cell(i, 5).fill = PatternFill("solid", fgColor=COLOR[color_c[f[4]]])
    wk.freeze_panes = "A2"
    wk.auto_filter.ref = f"A1:F{len(cambios) + 1}"

    # ---- Unidades ----
    wu = wb.create_sheet("Unidades")
    titulos = ["Censo", "Nivel", "Nombre impreso", "Región", "Provincia", "Departamento", "Partido", "Población",
               "Clave de comparación", "Cuadro"]
    cabecera(wu, 1, titulos, [8, 13, 30, 22, 20, 20, 14, 12, 26, 18])
    niv = {"region": 0, "provincia": 1, "partido": 1, "departamento": 2, "comuna": 3, "distrito": 3}
    uu = u.assign(o=u.nivel.map(niv)).sort_values(["censo", "o"], kind="stable")
    for i, f in enumerate(uu.itertuples(), 2):
        vals = [f.censo, f.nivel, f.nombre, f.region, f.provincia, f.departamento, f.partido,
                None if pd.isna(f.poblacion) or f.poblacion == "" else int(f.poblacion), f.clave, f.fuente]
        for j, v in enumerate(vals, 1):
            c = celda(wu, i, j, v)
            if j == 8:
                c.number_format = "#,##0"
    wu.freeze_panes = "A2"
    wu.auto_filter.ref = f"A1:J{len(uu) + 1}"
    n_u = len(uu) + 1

    # ---- Resumen (fórmulas sobre Unidades y Cambios) ----
    wr = wb.create_sheet("Resumen", 0)
    celda(wr, 1, 1, "Divisiones territoriales de los censos de Chile, 1813-1982", font=F_TITULO)
    celda(wr, 2, 1, "Conteos calculados sobre la hoja Unidades (unidades con cifra en el cuadro indicado) y la hoja "
          "Cambios (frente al censo anterior del mismo nivel).", font=F_NOTA)
    niveles = ["partido", "distrito", "region", "provincia", "departamento", "comuna"]
    titulos = ["Censo", "Partidos", "Distritos", "Regiones", "Provincias", "Departamentos", "Comunas",
               "Provincias nuevas o que vuelven", "Provincias que cambian de departamentos", "Provincias que dejan de aparecer",
               "Departamentos nuevos o que vuelven", "Departamentos que cambian de provincia", "Departamentos que dejan de aparecer",
               "Comunas nuevas (por nombre)", "Comunas que dejan de aparecer", "Cuadro fuente de la división",
               "Observaciones"]
    cabecera(wr, 4, titulos, [8, 9, 9, 9, 10, 13, 9, 11, 13, 12, 12, 13, 13, 12, 12, 48, 70])
    obs = {
        1813: "Partidos y distritos coloniales (cobertura parcial del censo de Egaña); sin equivalencia directa con los "
              "departamentos posteriores.",
        1835: "9 provincias y 50 departamentos según el propio censo. El sinóptico de 1865 agrupa estos departamentos en "
              "menos unidades (p. ej. los de Chiloé en Ancud, Castro, Quinchao y Carelmapu).",
        1843: "Solo departamentos y provincias (sinóptico de 1865). Atacama dentro de Coquimbo; Ñuble repartido entre "
              "Concepción y Maule.",
        1854: "Solo departamentos y provincias (sinóptico de 1865); misma agrupación que 1865.",
        1865: "Aparecen las provincias de Llanquihue, Arauco, Ñuble y Atacama ya en 1854; Curicó sigue en Colchagua.",
        1875: "Nuevas provincias de Curicó, Linares, Bío-Bío y el territorio de Angol.",
        1885: "Primer censo con Antofagasta, Tarapacá y Tacna; nueva provincia de O'Higgins.",
        1895: "Malleco y Cautín (antes territorio de Angol y Arauco); sin Tacna en el cuadro general.",
        1907: "Vuelve Tacna; nuevos departamentos de Santa Cruz, Chanco y Llaima.",
        1920: "Último censo con la división del siglo XIX; el cuadro 1 reconstruye 1865-1907 en estos límites.",
        1930: "Reforma de 1928: 17 provincias y 65 departamentos; primer censo con comunas. Aconcagua incluye "
              "Valparaíso y Colchagua a O'Higgins.",
        1940: "Se restablecen provincias (Valparaíso, O'Higgins, Curicó, Linares, Arauco, Malleco) y se crea Osorno: 25 "
              "provincias.",
        1952: "Mismas 25 provincias más el territorio antártico.",
        1960: "Mismas provincias (cuadro 1.1.1 de 1960 sin diferencias con 1940 y 1952). Departamentos no transcritos.",
        1970: "25 provincias; departamentos y comunas no transcritos.",
        1982: "Regionalización de 1974: 13 regiones, 51 provincias y comunas nuevas. Sin continuidad de nombre con "
              "las provincias anteriores.",
    }
    for i, c in enumerate(CENSOS, 5):
        celda(wr, i, 1, c, font=F_BOLD)
        for j, nv in enumerate(niveles, 2):
            celda(wr, i, j, f'=IF(COUNTIFS(Unidades!$A$2:$A${n_u},$A{i},Unidades!$B$2:$B${n_u},"{nv}")=0,"",'
                            f'COUNTIFS(Unidades!$A$2:$A${n_u},$A{i},Unidades!$B$2:$B${n_u},"{nv}"))', centro=True)
        n_k = len(cambios) + 1
        for j, (nv, tipos) in enumerate([("provincia", ("nueva", "vuelve a aparecer")),
                                         ("provincia", ("cambia departamentos",)), ("provincia", ("deja de aparecer",)),
                                         ("departamento", ("nuevo", "vuelve a aparecer")),
                                         ("departamento", ("cambia de provincia",)),
                                         ("departamento", ("deja de aparecer",)),
                                         ("comuna", ("nueva",)), ("comuna", ("deja de aparecer",))], 8):
            cuenta = "+".join(f'COUNTIFS(Cambios!$B$2:$B${n_k},$A{i},Cambios!$C$2:$C${n_k},"{nv}",'
                              f'Cambios!$E$2:$E${n_k},"{t}")' for t in tipos)
            celda(wr, i, j, f'=IF({cuenta}=0,"",{cuenta})', centro=True)
        celda(wr, i, 16, CUADRO_TEXTO[c], wrap=True)
        celda(wr, i, 17, obs[c], wrap=True)
    wr.freeze_panes = "B5"

    # ---- Léeme ----
    wl = wb.create_sheet("Léeme", 0)
    wl.column_dimensions["A"].width = 24
    wl.column_dimensions["B"].width = 110
    textos = [
        ("Divisiones territoriales de los censos de Chile (1813-1982)", None),
        ("", None),
        ("Qué es", "Las unidades territoriales con que cada censo publicó su población (provincia, departamento, comuna; "
                   "región desde 1982) y cómo se relacionan entre censos. Se arma con los cuadros ya transcritos del "
                   "proyecto, sin cifras escritas a mano. Generada por herramientas/divisiones_territoriales.py."),
        ("Resumen", "Una fila por censo: cuántas unidades tiene cada nivel y cuántas cambian frente al censo anterior."),
        ("Provincias", "Matriz provincia × censo (1835-1970): = mismos departamentos que el censo anterior, Δ cambian, "
                       "nueva. Al lado, los tramos sin cambios y los departamentos de cada censo."),
        ("Departamentos", "Matriz departamento × censo (1835-1952) con la provincia a la que pertenece en cada censo y sus "
                          "traslados."),
        ("Comunas", "Matriz comuna × censo (1930, 1940, 1952 y 1982), igualadas por nombre."),
        ("Cambios", "Lista de creaciones, supresiones y traslados entre censos consecutivos (filtrable)."),
        ("Comprobaciones", "Cifras que los cuadros comparativos dan para censos anteriores frente a la cifra original de "
                           "ese censo. Igual = la unidad tenía los mismos límites en ambos censos (prueba de "
                           "comparabilidad); distinta = cambió (o hay errata, ver las notas de los cuadros)."),
        ("Series comparables", "Qué cuadros comparativos existen, qué división usan y qué series permiten armar."),
        ("Unidades", "Lista larga: censo, nivel, nombre impreso, unidad superior, población total y cuadro de origen."),
        ("", None),
        ("Cómo se comparan", "Las unidades se igualan por nombre normalizado (sin tildes; variantes como Laja/La Laja, "
                             "Puchacai/Puchacay, Unión/La Unión, Aysen/Aisén). El mismo nombre no garantiza los mismos "
                             "límites: los departamentos cambian por dentro (subdelegaciones) sin cambiar de nombre. "
                             "La comparabilidad se prueba con los cuadros comparativos (hoja Comprobaciones)."),
        ("Advertencias", "1843 y 1854 no tienen publicación propia: vienen del sinóptico de 1865. 1960 y 1970 solo tienen "
                         "provincias transcritas. 1982 usa la DPA de 1974 (regiones, provincias y comunas nuevas), sin "
                         "equivalencia por nombre con las provincias anteriores. Tacna, Arica y Tarapacá: chilenas desde "
                         "1885 (Tacna sale después de 1929). Población: cifra impresa en el cuadro indicado (en 1940, "
                         "la de la publicación, no la revisada)."),
        ("Pendiente", "Departamentos de 1960 y 1970; comunas de 1960 y 1970; cuadro de cifras comparativas 1970-1982 por "
                      "comuna; leyes de creación de cada unidad (no incluidas)."),
    ]
    for i, (a, b) in enumerate(textos, 1):
        celda(wl, i, 1, a, font=F_TITULO if i == 1 else F_BOLD)
        if b:
            celda(wl, i, 2, b, wrap=True)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(SALIDA.relative_to(RAIZ), f"{len(u)} unidades, {len(cambios)} cambios, {len(comp_filas)} comprobaciones")


if __name__ == "__main__":
    main()
