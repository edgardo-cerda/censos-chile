"""Genera el diccionario consolidado de microdatos 1992–2024 en Excel.

Junta tres fuentes del repositorio (no lee los microdatos):
  - metadatos/inventario/        variables y frecuencias observadas en cada base,
  - metadatos/diccionario/       etiquetas de variables y categorías (oficiales o transcritas),
  - metadatos/concordancia/      conceptos y variable correspondiente en cada censo.

Salida: productos/diccionario/diccionario_censos.xlsx, con las hojas
  leeme, matriz (concepto × censo), comparacion (categorías lado a lado por concepto),
  variables (una fila por variable) y categorias (una fila por código).

Uso:
  python herramientas/generar_diccionario.py
"""

from __future__ import annotations

import csv
import fnmatch
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
INVENTARIO = RAIZ / "metadatos" / "inventario"
DICCIONARIO = RAIZ / "metadatos" / "diccionario"
CONCORDANCIA = RAIZ / "metadatos" / "concordancia" / "concordancia_microdatos.csv"
SALIDA = RAIZ / "productos" / "diccionario" / "diccionario_censos.xlsx"

CENSOS = ["1992", "2002", "2017", "2024"]
ORDEN_TABLAS = ["portafolios", "viviendas", "hogares", "uso_piezas", "personas"]
# Columnas agregadas por la exportación y que no son variables del censo.
IGNORAR = {"_fila"}
# En variables numéricas o territoriales solo se listan los códigos con etiqueta (especiales).
TIPOS_CON_TODOS_LOS_CODIGOS = {"categorica"}

FUENTES = {
    "diccionario_oficial": "Diccionario de variables publicado por el INE (2024).",
    "etiquetas_oficiales": "Archivos etiquetas_*.csv incluidos en los microdatos públicos (2017).",
    "tabla_auxiliar": "Tabla de códigos entregada junto con los microdatos.",
    "cuestionario": "Transcrito del cuestionario censal (código impreso junto a la opción).",
    "cuestionario_orden": "Cuestionario sin código impreso: se asume el orden de las opciones.",
    "inferido": "Deducido de la estructura de la base y las frecuencias; por verificar.",
    "entrega": "Variable técnica de la entrega (identificadores, geografía).",
    "derivada": "Variable construida por el INE a partir de otras preguntas.",
    "sin_documentacion": "Variable presente en la base sin documentación disponible.",
}
ESTADOS = {
    "documentado": "Código observado en la base y con etiqueta.",
    "sin etiqueta": "Código observado en la base sin etiqueta en las fuentes disponibles.",
    "no observado": "Código documentado que no aparece en la base.",
    "valor perdido": "Celda vacía o NA en el archivo original.",
}

AZUL = PatternFill("solid", fgColor="1F4E78")
GRIS = PatternFill("solid", fgColor="EDEDED")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
CELESTE = PatternFill("solid", fgColor="DDEBF7")
COLOR_CENSO = {"1992": "E2EFDA", "2002": "FCE4D6", "2017": "DDEBF7", "2024": "EDE7F6"}
BLANCO = Font(bold=True, color="FFFFFF")
NEGRITA = Font(bold=True)
BORDE = Border(bottom=Side(style="thin", color="BFBFBF"))


def leer_csv(ruta: Path) -> list[dict]:
    with ruta.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def normalizar(codigo: str) -> str:
    texto = (codigo or "").strip().strip("'")
    if re.fullmatch(r"-?\d+(\.0+)?", texto):
        return str(int(float(texto)))
    return texto


def clave_orden(codigo: str):
    """Números en orden numérico, luego textos, y al final los valores perdidos."""
    if codigo in ("", "NA"):
        return (2, 0, codigo)
    try:
        return (0, float(codigo), "")
    except ValueError:
        return (1, 0, codigo)


def cargar() -> tuple[list[dict], dict, list[dict]]:
    # Inventario
    variables = []
    for ruta in sorted(INVENTARIO.glob("*_variables.csv")):
        for r in leer_csv(ruta):
            if r["variable"] not in IGNORAR:
                variables.append(r)
    frecuencias = defaultdict(dict)
    for ruta in sorted(INVENTARIO.glob("*_frecuencias.csv")):
        for r in leer_csv(ruta):
            codigo = normalizar(r["valor"])
            if r["censo"] == "2024" and codigo == "":
                codigo = "NA"  # nulos del Parquet; el diccionario oficial los llama NA
            frecuencias[(r["censo"], r["tabla"], r["variable"])][codigo] = int(r["n"])

    # Diccionario
    doc_var = {}
    for ruta in sorted(DICCIONARIO.glob("*_variables.csv")):
        for r in leer_csv(ruta):
            doc_var[(r["censo"], r["tabla"], r["variable"])] = r
    etiquetas = defaultdict(dict)
    for ruta in sorted(DICCIONARIO.glob("*_categorias*.csv")):
        for r in leer_csv(ruta):
            etiquetas[(r["censo"], r["tabla"], r["variable"])][normalizar(r["codigo"])] = r

    for v in variables:
        clave = (v["censo"], v["tabla"], v["variable"])
        d = doc_var.get(clave, {})
        v["doc"] = d
        v["tipo"] = d.get("tipo") or ("categorica" if int(v["n_distintos"]) <= 30 else "numerica")
        v["frecuencias"] = frecuencias.get(clave, {})
        v["etiquetas"] = etiquetas.get(clave, {})
    variables.sort(key=lambda v: (v["censo"], ORDEN_TABLAS.index(v["tabla"]) if v["tabla"] in ORDEN_TABLAS else 99, int(v["orden"])))
    return variables, etiquetas, leer_csv(CONCORDANCIA)


def categorias_de(v: dict) -> list[dict]:
    """Une códigos observados y documentados de una variable."""
    obs, etq = v["frecuencias"], v["etiquetas"]
    codigos = set(etq)
    if v["tipo"] in TIPOS_CON_TODOS_LOS_CODIGOS:
        codigos |= set(obs)
    else:
        codigos |= {c for c in obs if c in ("", "NA")}
    total = int(v["n_filas"]) or 1
    filas = []
    for codigo in sorted(codigos, key=clave_orden):
        e = etq.get(codigo, {})
        etiqueta, fuente, nota = e.get("etiqueta", ""), e.get("fuente", ""), e.get("nota", "")
        if codigo == "NA" and v["censo"] == "2024" and "NA (no aplica)" in v["doc"].get("nota", ""):
            etiqueta, fuente = "No aplica", "diccionario_oficial"
        n = obs.get(codigo)
        if n is None and fuente == "tabla_auxiliar":
            continue  # las tablas de códigos traen niveles y códigos que no se usan en la base
        if codigo in ("", "NA") and not etiqueta:
            estado = "valor perdido"
        elif n is None:
            estado = "no observado" if v["frecuencias"] else ""
        else:
            estado = "documentado" if etiqueta else "sin etiqueta"
        filas.append({"codigo": codigo if codigo else "(vacío)", "etiqueta": etiqueta, "n": n,
                      "pct": (n / total) if n is not None else None, "estado": estado, "fuente": fuente, "nota": nota})
    return filas


def conceptos_por_variable(concordancia: list[dict], variables: list[dict]) -> tuple[dict, dict]:
    """Resuelve las celdas de la matriz de concordancia a variables reales de cada censo."""
    nombres = defaultdict(set)
    for v in variables:
        nombres[v["censo"]].add(v["variable"])
    por_variable = defaultdict(list)
    por_concepto = {}
    for fila in concordancia:
        por_concepto[fila["concepto"]] = {}
        for censo in CENSOS:
            celda = re.sub(r"\(.*?\)", "", fila.get(f"v{censo}", ""))
            encontrados = []
            for token in re.split(r"[;+/→,…\s]+", celda):
                token = re.sub(r"^[A-Z_]+\.", "", token.strip())
                if not token:
                    continue
                if "*" in token:
                    encontrados += sorted(fnmatch.filter(nombres[censo], token))
                elif token in nombres[censo]:
                    encontrados.append(token)
            por_concepto[fila["concepto"]][censo] = encontrados
            for nombre in encontrados:
                por_variable[(censo, nombre)].append(fila["concepto"])
    return por_variable, por_concepto


# ----------------------------------------------------------------------------- Excel

def encabezado(ws, columnas: list[tuple[str, int]], fila: int = 1) -> None:
    for i, (titulo, ancho) in enumerate(columnas, 1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font, c.fill = BLANCO, AZUL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def hoja_leeme(wb: Workbook, variables: list[dict], n_categorias: int) -> None:
    ws = wb.active
    ws.title = "leeme"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    lineas = [
        ("Diccionario consolidado de microdatos censales 1992–2024", None),
        ("", None),
        ("Qué contiene", "Variables y categorías de respuesta de las bases de microdatos de los censos 1992, 2002, 2017 y 2024, "
                         "con el nombre de cada variable en la base, su pregunta, universo y los códigos con su etiqueta y frecuencia."),
        ("Generado por", "herramientas/generar_diccionario.py a partir de metadatos/inventario, metadatos/diccionario y "
                         "metadatos/concordancia. Para corregir una etiqueta, editar el CSV en metadatos/diccionario y volver a generar."),
        ("Totales", f"{len(variables)} variables y {n_categorias} códigos."),
        ("", None),
        ("Hojas", None),
        ("matriz", "Concepto × censo: qué variable de cada base corresponde a cada concepto (— = no existe)."),
        ("comparacion", "Para cada concepto, las categorías de cada censo lado a lado, con código, etiqueta y frecuencia."),
        ("variables", "Una fila por variable de cada tabla: etiqueta, pregunta, universo, tipo, conceptos y estadísticas básicas."),
        ("categorias", "Una fila por código de cada variable: etiqueta, frecuencia, porcentaje, estado y fuente de la etiqueta."),
        ("", None),
        ("Fuente de la etiqueta", None),
        *FUENTES.items(),
        ("", None),
        ("Estado del código", None),
        *ESTADOS.items(),
        ("", None),
        ("Notas", "Los microdatos de 1982 no se incluyen (ver docs/notas_microdatos.md). En variables numéricas y territoriales "
                  "solo se listan los códigos especiales documentados; los códigos territoriales van en metadatos/territorio. "
                  "Las frecuencias corresponden a registros de cada tabla (personas, viviendas u hogares)."),
    ]
    for i, (a, b) in enumerate(lineas, 1):
        ws.cell(row=i, column=1, value=a)
        if b is None:
            ws.cell(row=i, column=1).font = Font(bold=True, size=14 if i == 1 else 11)
        else:
            ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=i, column=1).alignment = Alignment(vertical="top")


def hoja_matriz(wb: Workbook, concordancia: list[dict], por_concepto: dict, variables: list[dict]) -> None:
    ws = wb.create_sheet("matriz")
    etiqueta = {(v["censo"], v["variable"]): v["doc"].get("etiqueta", "") for v in variables}
    columnas = [("Tema", 13), ("Concepto", 24), ("Unidad", 10)] + [(c, 34) for c in CENSOS] + [("Estado", 10), ("Nota", 60)]
    encabezado(ws, columnas)
    for i, fila in enumerate(concordancia, 2):
        ws.cell(row=i, column=1, value=fila["tema"])
        ws.cell(row=i, column=2, value=fila["concepto"]).font = NEGRITA
        ws.cell(row=i, column=3, value=fila["unidad"])
        for j, censo in enumerate(CENSOS, 4):
            texto = fila[f"v{censo}"]
            vars_ = por_concepto[fila["concepto"]][censo]
            if len(vars_) == 1 and etiqueta.get((censo, vars_[0])):
                texto = f"{texto}\n{etiqueta[(censo, vars_[0])]}"
            c = ws.cell(row=i, column=j, value=texto)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = GRIS if texto.strip() in ("—", "") else PatternFill("solid", fgColor=COLOR_CENSO[censo])
        ws.cell(row=i, column=8, value=fila["estado"])
        ws.cell(row=i, column=9, value=fila["nota"]).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(concordancia) + 1}"


def hoja_comparacion(wb: Workbook, concordancia: list[dict], por_concepto: dict, variables: list[dict]) -> None:
    ws = wb.create_sheet("comparacion")
    indice = {(v["censo"], v["variable"]): v for v in variables}
    columnas = [("Tema", 12), ("Concepto", 22)]
    for censo in CENSOS:
        columnas += [(f"{censo} código", 12), (f"{censo} etiqueta", 34), (f"{censo} n", 11)]
    encabezado(ws, columnas)
    fila = 2
    for concepto in concordancia:
        bloques = {}
        for censo in CENSOS:
            # Variable principal: la primera categórica con categorías documentadas.
            candidatas = [indice[(censo, n)] for n in por_concepto[concepto["concepto"]][censo] if (censo, n) in indice]
            candidatas = [v for v in candidatas if v["tipo"] == "categorica" and v["etiquetas"]]
            if candidatas:
                v = candidatas[0]
                bloques[censo] = (v, [c for c in categorias_de(v) if c["estado"] != "valor perdido"])
        if not bloques:
            continue
        ws.cell(row=fila, column=1, value=concepto["tema"]).font = NEGRITA
        ws.cell(row=fila, column=2, value=concepto["concepto"]).font = NEGRITA
        for k, censo in enumerate(CENSOS):
            col = 3 + 3 * k
            for c in range(col, col + 3):
                ws.cell(row=fila, column=c).fill = PatternFill("solid", fgColor=COLOR_CENSO[censo])
            if censo in bloques:
                v = bloques[censo][0]
                ws.cell(row=fila, column=col, value=f'{v["variable"]} [{v["tabla"]}]').font = NEGRITA
                ws.cell(row=fila, column=col + 1, value=v["doc"].get("etiqueta", "")).font = NEGRITA
            else:
                ws.cell(row=fila, column=col + 1, value="—")
        for c in range(1, len(columnas) + 1):
            ws.cell(row=fila, column=c).border = BORDE
        largo = max(len(b[1]) for b in bloques.values())
        for i in range(largo):
            for k, censo in enumerate(CENSOS):
                if censo not in bloques or i >= len(bloques[censo][1]):
                    continue
                cat = bloques[censo][1][i]
                col = 3 + 3 * k
                ws.cell(row=fila + 1 + i, column=col, value=cat["codigo"])
                c = ws.cell(row=fila + 1 + i, column=col + 1, value=cat["etiqueta"] or f"[{cat['estado']}]")
                if cat["estado"] != "documentado":
                    c.fill = AMARILLO if cat["estado"] == "sin etiqueta" else GRIS
                ws.cell(row=fila + 1 + i, column=col + 2, value=cat["n"]).number_format = "#,##0"
        fila += largo + 2


def hoja_variables(wb: Workbook, variables: list[dict], por_variable: dict) -> None:
    ws = wb.create_sheet("variables")
    columnas = [("Censo", 7), ("Tabla", 11), ("Orden", 6), ("Variable", 22), ("Etiqueta", 40), ("Pregunta", 50), ("Universo", 28),
                ("Tipo", 12), ("Fuente", 18), ("Conceptos", 20), ("Registros", 12), ("Valores distintos", 10),
                ("% vacío o NA", 9), ("Mínimo", 9), ("Máximo", 9), ("Códigos sin etiqueta", 10), ("Nota", 50)]
    encabezado(ws, columnas)
    for i, v in enumerate(variables, 2):
        d = v["doc"]
        cats = categorias_de(v)
        sin_etiqueta = sum(c["estado"] == "sin etiqueta" for c in cats)
        n = int(v["n_filas"]) or 1
        valores = [int(v["censo"]), v["tabla"], int(v["orden"]), v["variable"], d.get("etiqueta", ""), d.get("pregunta", ""),
                   d.get("universo", ""), v["tipo"], d.get("fuente", "sin_documentacion"),
                   "; ".join(por_variable.get((v["censo"], v["variable"]), [])), int(v["n_filas"]), int(v["n_distintos"]),
                   (int(v["n_vacios"]) + int(v["n_na"])) / n, v["min"], v["max"], sin_etiqueta or None, d.get("nota", "")]
        for j, valor in enumerate(valores, 1):
            c = ws.cell(row=i, column=j, value=valor)
            if j in (5, 6, 7, 17):
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=COLOR_CENSO[v["censo"]])
        ws.cell(row=i, column=4).font = NEGRITA
        ws.cell(row=i, column=11).number_format = "#,##0"
        ws.cell(row=i, column=13).number_format = "0.0%"
        if sin_etiqueta:
            ws.cell(row=i, column=16).fill = AMARILLO
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(variables) + 1}"


def hoja_categorias(wb: Workbook, variables: list[dict]) -> int:
    ws = wb.create_sheet("categorias")
    columnas = [("Censo", 7), ("Tabla", 11), ("Variable", 22), ("Etiqueta de la variable", 36), ("Código", 9),
                ("Etiqueta del código", 48), ("n", 12), ("%", 8), ("Estado", 13), ("Fuente", 18), ("Nota", 50)]
    encabezado(ws, columnas)
    fila = 2
    for v in variables:
        for cat in categorias_de(v):
            valores = [int(v["censo"]), v["tabla"], v["variable"], v["doc"].get("etiqueta", ""), cat["codigo"], cat["etiqueta"],
                       cat["n"], cat["pct"], cat["estado"], cat["fuente"], cat["nota"]]
            for j, valor in enumerate(valores, 1):
                ws.cell(row=fila, column=j, value=valor)
            ws.cell(row=fila, column=1).fill = PatternFill("solid", fgColor=COLOR_CENSO[v["censo"]])
            ws.cell(row=fila, column=7).number_format = "#,##0"
            ws.cell(row=fila, column=8).number_format = "0.0%"
            if cat["estado"] == "sin etiqueta":
                ws.cell(row=fila, column=9).fill = AMARILLO
            elif cat["estado"] in ("no observado", "valor perdido"):
                ws.cell(row=fila, column=9).fill = GRIS
            fila += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{fila - 1}"
    return fila - 2


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    variables, _, concordancia = cargar()
    por_variable, por_concepto = conceptos_por_variable(concordancia, variables)
    wb = Workbook()
    # La primera hoja (leeme) se llena al final, cuando se conocen los totales.
    hoja_matriz(wb, concordancia, por_concepto, variables)
    hoja_comparacion(wb, concordancia, por_concepto, variables)
    hoja_variables(wb, variables, por_variable)
    n_categorias = hoja_categorias(wb, variables)
    hoja_leeme(wb, variables, n_categorias)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)

    sin_etiqueta = sum(1 for v in variables for c in categorias_de(v) if c["estado"] == "sin etiqueta")
    print(f"{SALIDA.relative_to(RAIZ)}: {len(variables)} variables, {n_categorias} códigos "
          f"({sin_etiqueta} observados sin etiqueta).")


if __name__ == "__main__":
    main()
